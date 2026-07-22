import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from jsonschema import Draft202012Validator
from openpyxl import Workbook
import pytest

from core.controlling import artifacts as artifacts_module
import core.controlling.tools as controlling_tools
from core.controlling.tools import ControllingComputeTool, InstanceReadTool
from core.controlling.narrative import ControllingReportTool


def _write_compute_mapping(
    tmp_path, groups, *, source_globs=None, threshold=1
):
    mapping = _strict_mapping(
        groups=groups,
        metrics={"ebitda": [next(iter(groups))]},
        row_identifier="Konto",
    )
    mapping["thresholds"] = {"material_deviation_abs": threshold}
    if source_globs is not None:
        mapping["source_globs"] = source_globs
    pilot = tmp_path / "project" / "pilot"
    pilot.mkdir(parents=True)
    (pilot / "account-mapping.json").write_text(
        json.dumps(mapping), encoding="utf-8"
    )


def _strict_mapping(*, groups=None, metrics=None, row_identifier="Zeilen-ID"):
    groups = groups or {
        "revenue": {"prefixes": ["4"], "multiplier": 1},
        "liquidity": {"prefixes": ["1"], "multiplier": 1},
    }
    groups = {
        name: {**rule, "multiplier": rule.get("multiplier", 1)}
        for name, rule in groups.items()
    }
    metrics = dict(metrics or {"ebitda": ["revenue"], "egt": ["revenue"]})
    first_group = next(iter(groups))
    metrics.setdefault("ebitda", [first_group])
    metrics.setdefault("depreciation", [first_group])
    metrics.setdefault("egt", [first_group])
    return {
        "schema_version": 1,
        "report_metadata": {
            "period_start": "2026-06-01",
            "period_end": "2026-06-30",
            "currency": "EUR",
            "sign_convention": "source_value_times_group_multiplier",
            "row_identifier": row_identifier,
        },
        "columns": {
            "account_id": "Konto",
            "account_name": "Bezeichnung",
            "month": "Monat",
            "prior": "Vormonat",
            "ytd": "YTD",
            "ytd_prior": "YTD Vorjahr",
        },
        "sign_semantics": {
            "formula": "normalized_value = source_value * group.multiplier",
            "positive_contribution": "Positive increases KPI.",
            "negative_contribution": "Negative decreases KPI.",
        },
        "source_globs": ["data/*.csv"],
        "groups": groups,
        "metrics": metrics,
        "thresholds": {"material_deviation_abs": 1},
    }


def _write_strict_mapping(tmp_path, **overrides):
    mapping = _strict_mapping()
    mapping.update(overrides)
    pilot = tmp_path / "project" / "pilot"
    pilot.mkdir(parents=True, exist_ok=True)
    (pilot / "account-mapping.json").write_text(
        json.dumps(mapping), encoding="utf-8"
    )
    return mapping


def test_read_lists_and_reads_only_approved_instance_inputs(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "bwa.csv").write_text(
        "konto,wert\nUmsatz,120\n", encoding="utf-8"
    )
    tool = InstanceReadTool(tmp_path)

    metadata = tool.safe_run("read:data/bwa.csv")
    assert "SOURCE: data/bwa.csv" in metadata
    assert "ROWS: 1" in metadata
    assert "Umsatz,120" not in metadata
    assert "data/bwa.csv" in tool.safe_run("list")
    assert tool.safe_run("read:../secret.csv").startswith("FEHLER")


def test_compute_is_deterministic_and_cites_csv_source(tmp_path):
    (tmp_path / "imports").mkdir()
    (tmp_path / "imports" / "plan.csv").write_text(
        "ist,budget\n120,100\n80,100\n", encoding="utf-8"
    )
    tool = ControllingComputeTool(tmp_path)

    result = tool.safe_run(
        json.dumps(
            {
                "operation": "difference",
                "source": "imports/plan.csv",
                "actual_column": "ist",
                "budget_column": "budget",
            }
        )
    )

    assert "RESULT: 0" in result
    assert "SOURCE: imports/plan.csv" in result
    assert "METHOD: sum(ist) - sum(budget)" in result


def test_compute_fails_closed_for_outside_file_and_invalid_numbers(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "bad.csv").write_text(
        "value\nnot-a-number\n", encoding="utf-8"
    )
    tool = ControllingComputeTool(tmp_path)

    assert tool.safe_run(
        '{"operation":"sum","source":"../outside.csv","column":"value"}'
    ).startswith("FEHLER")
    assert tool.safe_run(
        '{"operation":"sum","source":"data/bad.csv","column":"value"}'
    ).startswith("FEHLER")


def test_compute_preserves_decimal_separator_meaning(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "amounts.csv").write_text(
        "value;other\n1.200,50;x\n1,25;y\n", encoding="utf-8"
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/amounts.csv","column":"value"}'
    )

    assert "RESULT: 1201.75" in result


@pytest.mark.parametrize("raw", ["1.234", "1,234"])
def test_compute_rejects_ambiguous_single_separator_numbers(tmp_path, raw):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "ambiguous.csv").write_text(
        f"value;note\n{raw};x\n", encoding="utf-8"
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/ambiguous.csv","column":"value"}'
    )

    assert result.startswith("FEHLER")


@pytest.mark.parametrize("raw", ["1.234,56", "1,234.56"])
def test_compute_accepts_unambiguous_grouped_decimals(tmp_path, raw):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "amount.csv").write_text(
        f"value;note\n{raw};x\n", encoding="utf-8"
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/amount.csv","column":"value"}'
    )

    assert "RESULT: 1234.56" in result


def test_compute_rejects_oversized_source_before_parsing(tmp_path):
    (tmp_path / "data").mkdir()
    source = tmp_path / "data" / "oversized.csv"
    with source.open("wb") as handle:
        handle.seek(25 * 1024 * 1024)
        handle.write(b"\n")

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/oversized.csv","column":"value"}'
    )

    assert result.startswith("FEHLER: Eingabedatei ist zu gro")


def test_compute_rejects_more_than_100_000_rows(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "too-many.csv").write_text(
        "value\n" + ("1\n" * 100_001), encoding="utf-8"
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/too-many.csv","column":"value"}'
    )

    assert result.startswith("FEHLER: Eingabe enth\u00e4lt zu viele Zeilen")


def test_kmu_report_uses_explicit_mapping_and_returns_only_aggregates(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    (tmp_path / "data" / "salden.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        "4000;Umsatz Inland;1000;900;2900;2700\n"
        "5000;Material;300;250;800;750\n"
        "6000;Lohn;200;180;560;520\n"
        "1200;Kunde Alpha;450;400;450;\n"
        "1000;Bank;600;650;600;\n"
        "1600;Lieferant Beta;-350;-300;-350;\n",
        encoding="utf-8",
    )
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(
                groups={
                    "revenue": {"prefixes": ["4"]},
                    "material": {"prefixes": ["5"]},
                    "personnel": {"prefixes": ["6"]},
                    "receivables": {"prefixes": ["12"]},
                    "liquidity": {"prefixes": ["10"]},
                    "payables": {"prefixes": ["16"]},
                },
                metrics={
                    "ebitda": ["revenue", "material", "personnel"],
                    "egt": ["revenue", "material", "personnel"],
                },
                row_identifier="Konto",
            )
            | {"thresholds": {"material_deviation_abs": 40}}
        ),
        encoding="utf-8",
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/salden.csv"}'
    )
    report = json.loads(result)

    assert report["metrics"]["ebitda"]["month"] == "1500"
    assert report["metrics"]["material_ratio"]["month"] == "0.3"
    assert report["metrics"]["receivables"]["month"] == "450"
    assert any(
        item["account_id"] == "5000" for item in report["material_deviations"]
    )
    assert "Kunde Alpha" not in result


def test_kmu_report_fails_closed_without_customer_mapping(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "salden.csv").write_text(
        "Konto;Monat\n4000;100\n", encoding="utf-8"
    )

    assert (
        ControllingComputeTool(tmp_path)
        .safe_run('{"operation":"kmu_report","source":"data/salden.csv"}')
        .startswith("FEHLER")
    )


def test_kmu_report_auto_source_requires_exactly_one_input(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "imports").mkdir()
    (tmp_path / "data" / "one.csv").write_text("x\n1\n", encoding="utf-8")
    (tmp_path / "imports" / "two.csv").write_text("x\n2\n", encoding="utf-8")

    assert (
        ControllingComputeTool(tmp_path)
        .safe_run('{"operation":"kmu_report","source":"auto"}')
        .startswith("FEHLER")
    )


def test_kmu_report_auto_source_uses_mapping_glob_not_optional_inputs(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "imports").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    (tmp_path / "data" / "actuals.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        "4000;Umsatz;10;9;10;9\n",
        encoding="utf-8",
    )
    (tmp_path / "imports" / "budget.csv").write_text("unused\n1\n", encoding="utf-8")
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(row_identifier="Konto")
            | {"source_globs": ["data/actuals*.csv"]}
        ),
        encoding="utf-8",
    )

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"auto"}'
        )
    )

    assert report["source"] == "data/actuals.csv"


def test_kmu_report_auto_source_rejects_symlink_resolving_outside_inputs(
    tmp_path, monkeypatch
):
    (tmp_path / "data").mkdir()
    outside = tmp_path.parent / f"{tmp_path.name}-outside.csv"
    outside.write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD\n4000;Umsatz;10;9;10\n",
        encoding="utf-8",
    )
    linked_source = tmp_path / "data" / "actuals.csv"
    try:
        linked_source.symlink_to(outside)
    except OSError:
        monkeypatch.setattr(
            type(tmp_path),
            "glob",
            lambda _root, _pattern: iter([outside]),
        )
    _write_compute_mapping(
        tmp_path,
        {"revenue": {"prefixes": ["4"]}},
        source_globs=["data/actuals.csv"],
    )
    tool = ControllingComputeTool(tmp_path)
    read_paths = []
    monkeypatch.setattr(
        tool,
        "_kmu_report",
        lambda path: read_paths.append(path) or "{}",
    )

    result = tool.safe_run('{"operation":"kmu_report","source":"auto"}')

    assert result.startswith("FEHLER")
    assert read_paths == []


@pytest.mark.parametrize(
    "rule",
    [
        {"prefixes": "4"},
        {"prefixes": [4]},
        {"prefixes": ["4"], "ignored": "yes"},
    ],
    ids=("string-prefixes", "non-string-prefix", "non-boolean-ignored"),
)
def test_kmu_report_returns_error_for_malformed_group_rules(tmp_path, rule):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "actuals.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD\n4000;Umsatz;10;9;10\n",
        encoding="utf-8",
    )
    _write_compute_mapping(tmp_path, {"revenue": rule})

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")


@pytest.mark.parametrize("multiplier", ["not-a-number", "NaN", None])
def test_kmu_report_returns_error_for_malformed_multiplier(tmp_path, multiplier):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "actuals.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD\n4000;Umsatz;10;9;10\n",
        encoding="utf-8",
    )
    _write_compute_mapping(
        tmp_path,
        {"revenue": {"prefixes": ["4"], "multiplier": multiplier}},
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")


@pytest.mark.parametrize("ignored", (False, True), ids=("included", "ignored"))
@pytest.mark.parametrize("multiplier", (0, -0, 0.0, "0", "-0.0"))
def test_kmu_report_rejects_zero_multiplier_and_requires_ignored_flag(
    tmp_path, ignored, multiplier
):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["groups"]["revenue"].update(multiplier=multiplier, ignored=ignored)
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "ignored:true" in result


@pytest.mark.parametrize("ignored", (False, True), ids=("included", "ignored"))
@pytest.mark.parametrize(
    "multiplier", (2, -2, 0.5, -0.5, "2", "-0.5", "1", "-1", True, False)
)
def test_kmu_report_rejects_non_sign_multiplier(tmp_path, ignored, multiplier):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["groups"]["revenue"].update(multiplier=multiplier, ignored=ignored)
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "vorzeichen" in result.casefold()


@pytest.mark.parametrize("multiplier", (1, -1, 1.0, -1.0))
def test_kmu_report_accepts_sign_only_multiplier(tmp_path, multiplier):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["groups"]["revenue"]["multiplier"] = multiplier
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert not result.startswith("FEHLER")


@pytest.mark.parametrize("multiplier", (0, -0, 0.0))
def test_mapping_schema_rejects_numeric_zero_multiplier(multiplier):
    schema = json.loads(
        (Path(__file__).parents[1] / "presets" / "kmu-controlling" / "pilot" / "account-mapping.schema.json").read_text(encoding="utf-8")
    )
    mapping = _strict_mapping()
    mapping["groups"]["revenue"]["multiplier"] = multiplier

    assert list(Draft202012Validator(schema).iter_errors(mapping))


@pytest.mark.parametrize("multiplier", (2, -2, 0.5, -0.5))
def test_mapping_schema_rejects_non_sign_multiplier(multiplier):
    schema = json.loads(
        (Path(__file__).parents[1] / "presets" / "kmu-controlling" / "pilot" / "account-mapping.schema.json").read_text(encoding="utf-8")
    )
    mapping = _strict_mapping()
    mapping["groups"]["revenue"]["multiplier"] = multiplier

    assert list(Draft202012Validator(schema).iter_errors(mapping))


def test_kmu_report_returns_error_for_malformed_source_glob(tmp_path):
    (tmp_path / "data").mkdir()
    _write_compute_mapping(
        tmp_path,
        {"revenue": {"prefixes": ["4"]}},
        source_globs=["data/**invalid/*.csv"],
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"auto"}'
    )

    assert result.startswith("FEHLER")


def test_kmu_report_deviations_cover_every_mapped_non_ignored_account(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "actuals.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        "4000;Umsatz;200;100;200;100\n"
        "5000;Kosten;40;100;40;100\n"
        "9000;Ignoriert;1000;0;1000;0\n",
        encoding="utf-8",
    )
    _write_compute_mapping(
        tmp_path,
        {
            "revenue": {"prefixes": ["4"]},
            "costs": {"prefixes": ["5"], "multiplier": -1},
            "ignored": {"prefixes": ["9"], "ignored": True},
        },
        threshold=50,
    )

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/actuals.csv"}'
        )
    )

    assert {
        item["account_id"]: item["month_vs_prior"]
        for item in report["material_deviations"]
    } == {"4000": "100", "5000": "60"}
    assert "ignored" not in report["groups"]


def test_kmu_report_fails_closed_for_unmapped_account(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    (tmp_path / "data" / "salden.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        "4000;Umsatz;10;9;10;9\n5000;Kosten;2;1;2;1\n",
        encoding="utf-8",
    )
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(
                groups={"revenue": {"prefixes": ["4"]}},
                metrics={"ebitda": ["revenue"]},
                row_identifier="Konto",
            )
        ),
        encoding="utf-8",
    )

    assert "Konto '5000' ist nicht zugeordnet" in ControllingComputeTool(
        tmp_path
    ).safe_run('{"operation":"kmu_report","source":"data/salden.csv"}')


def test_kmu_report_fails_closed_for_overlapping_mapping_rules(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    (tmp_path / "data" / "salden.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        "4000;Umsatz;10;9;10;9\n",
        encoding="utf-8",
    )
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(
                groups={
                    "revenue": {"prefixes": ["4"]},
                    "also_revenue": {"name_matches": ["umsatz"]},
                },
                metrics={"ebitda": ["revenue"]},
                row_identifier="Konto",
            )
        ),
        encoding="utf-8",
    )

    assert "Konto '4000' ist mehrfach zugeordnet" in ControllingComputeTool(
        tmp_path
    ).safe_run('{"operation":"kmu_report","source":"data/salden.csv"}')


def test_kmu_report_rejects_duplicate_semantic_column_assignments(tmp_path):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["columns"]["prior"] = mapping["columns"]["month"]
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "eindeutig" in result


def test_kmu_report_rejects_metrics_that_reference_ignored_groups(tmp_path):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping(
        groups={
            "revenue": {"prefixes": ["4"], "multiplier": 1},
            "excluded": {"prefixes": ["9"], "multiplier": 1, "ignored": True},
        },
        metrics={"ebitda": ["excluded"], "egt": ["revenue"]},
    )
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "ignorierte" in result


@pytest.mark.parametrize(
    "reserved_name",
    (
        "material_ratio",
        "personnel_ratio",
        "receivables",
        "liquidity",
        "payables",
    ),
)
def test_kmu_report_rejects_reserved_derived_metric_names(tmp_path, reserved_name):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["metrics"][reserved_name] = ["revenue"]
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "reserv" in result.casefold()


@pytest.mark.parametrize(
    "reserved_name",
    (
        "material_ratio",
        "personnel_ratio",
        "receivables",
        "liquidity",
        "payables",
    ),
)
def test_mapping_schema_rejects_reserved_derived_metric_names(reserved_name):
    schema = json.loads(
        (
            Path(__file__).parents[1]
            / "presets"
            / "kmu-controlling"
            / "pilot"
            / "account-mapping.schema.json"
        ).read_text(encoding="utf-8")
    )
    mapping = _strict_mapping()
    mapping["metrics"][reserved_name] = ["revenue"]

    assert list(Draft202012Validator(schema).iter_errors(mapping))


@pytest.mark.parametrize(
    ("mutation", "label"),
    [
        (lambda mapping: mapping.pop("source_globs"), "source_globs"),
        (lambda mapping: mapping["metrics"].pop("egt"), "Pflichtkennzahlen"),
        (
            lambda mapping: mapping["groups"]["revenue"].pop("multiplier"),
            "multiplier",
        ),
    ],
    ids=("source-globs", "required-metrics", "group-multiplier"),
)
def test_kmu_report_enforces_published_mapping_contract(tmp_path, mutation, label):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mutation(mapping)
    pilot = tmp_path / "project" / "pilot"
    pilot.mkdir(parents=True, exist_ok=True)
    (pilot / "account-mapping.json").write_text(
        json.dumps(mapping), encoding="utf-8"
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert label in result


def test_kmu_report_provenance_matches_exact_source_and_mapping_bytes(tmp_path):
    import hashlib

    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    _write_strict_mapping(tmp_path)

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/actuals.csv"}'
        )
    )

    source = tmp_path / "data" / "actuals.csv"
    mapping = tmp_path / "project" / "pilot" / "account-mapping.json"
    assert report["provenance"] == {
        "input_fingerprint": hashlib.sha256(source.read_bytes()).hexdigest(),
        "mapping_version": hashlib.sha256(mapping.read_bytes()).hexdigest(),
        "calculation_version": "kmu-report-v2",
    }


def test_kmu_report_finds_salden_header_after_workbook_title_rows(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Saldenliste März"
    sheet.append(["Saldenliste", None, None, None, None, None])
    sheet.append(["Stand 31.03.", None, None, None, None, None])
    sheet.append(
        ["Konto", "Bezeichnung", "Monat", "Vormonat", "YTD", "YTD Vorjahr"]
    )
    sheet.append(["4000", "Umsatz", 100, 90, 280, 270])
    workbook.save(tmp_path / "data" / "salden.xlsx")
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(
                groups={"revenue": {"prefixes": ["4"]}},
                metrics={"ebitda": ["revenue"], "egt": ["revenue"]},
                row_identifier="Konto",
            )
        ),
        encoding="utf-8",
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/salden.xlsx"}'
    )

    assert json.loads(result)["metrics"]["ebitda"]["ytd"] == "280"


def test_kmu_report_handles_realistic_two_section_saldenliste(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Saldenliste März"])
    sheet.append(["Mandant"])
    sheet.append(["Auswertung"])
    sheet.append(["Konto", "Bezeichnung", "Monat", "Vormonat", "YTD", "YTD Vorjahr"])
    sheet.append(["4000", "Umsatz", 203500, 200000, 600000, 550000])
    sheet.append(["5000", "Material", 101750, 100000, 300000, 280000])
    sheet.append(["6000", "Personal", 98950, 97000, 290000, 275000])
    sheet.append(["7000", "Abschreibung", 3000, 2900, 9000, 8500])
    sheet.append(["8000", "Finanzergebnis", 2050, 2000, 6000, 5500])
    sheet.append(["", "Zwischensumme", 0, 0, 0, 0])
    while sheet.max_row < 21:
        sheet.append([None] * 6)
    sheet.append(["Konto", "Bezeichnung", "Monat", "Vormonat", "YTD", "YTD Vorjahr"])
    sheet.append(["1200", "Forderungen", 238500, 230000, None, None])
    sheet.append(["1000", "Bank", 66150, 65000, None, None])
    workbook.save(tmp_path / "data" / "salden.xlsx")
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(
                groups={
                    "revenue": {"prefixes": ["4"]},
                    "material": {"prefixes": ["5"], "multiplier": -1},
                    "personnel": {"prefixes": ["6"], "multiplier": -1},
                    "depreciation": {"prefixes": ["7"], "multiplier": -1},
                    "financial": {"prefixes": ["8"], "multiplier": -1},
                    "receivables": {"prefixes": ["12"]},
                    "liquidity": {"prefixes": ["10"]},
                },
                metrics={
                    "ebitda": ["revenue", "material", "personnel"],
                    "egt": [
                        "revenue",
                        "material",
                        "personnel",
                        "depreciation",
                        "financial",
                    ],
                },
                row_identifier="Konto",
            )
            | {"thresholds": {"material_deviation_abs": 1000}}
        ),
        encoding="utf-8",
    )

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/salden.xlsx"}'
        )
    )

    assert report["groups"]["revenue"]["month"] == "203500"
    assert report["metrics"]["ebitda"]["month"] == "2800"
    assert report["metrics"]["egt"]["month"] == "-2250"
    assert report["metrics"]["material_ratio"]["month"] == "0.5"
    assert report["metrics"]["receivables"]["month"] == "238500"
    assert report["metrics"]["receivables"]["ytd"] is None
    assert report["metrics"]["liquidity"]["month"] == "66150"
    assert report["reconciliation"]["source_account_count"] == 7
    assert report["reconciliation"]["mapped_account_count"] == 7
    assert (
        report["reconciliation"]["mapped_source_period_totals"]
        == report["reconciliation"]["grouped_period_totals"]
    )


@pytest.mark.parametrize(
    "untrusted_response",
    [
        "Das EBITDA beträgt 8. Der YTD-Wert beträgt 10. Die Lage ist stabil.",
        "Das EBITDA beträgt acht. Die Lage ist stabil. Weitere Prüfung folgt.",
        "Die Liquidität beträgt ½. Die Lage ist stabil. Weitere Prüfung folgt.",
        "Es besteht ein qualitatives Liquiditätsrisiko. Die Lage ist angespannt. Prüfung folgt.",
    ],
    ids=("digit-swapped", "spelled-number", "unicode-number", "qualitative-risk"),
)
def test_narrative_is_deterministic_when_model_is_configured(
    tmp_path, monkeypatch, untrusted_response
):
    (tmp_path / "data").mkdir()
    (tmp_path / "project" / "pilot").mkdir(parents=True)
    (tmp_path / "data" / "actuals.csv").write_text(
        "Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        "4000;Geheimer Kunde;10;9;10;9\n",
        encoding="utf-8",
    )
    (tmp_path / "project" / "manifest.json").write_text(
        '{"model":{"narrative":"cheap"}}', encoding="utf-8"
    )
    (tmp_path / "project" / "pilot" / "account-mapping.json").write_text(
        json.dumps(
            _strict_mapping(
                groups={"revenue": {"prefixes": ["4"]}},
                metrics={"ebitda": ["revenue"], "egt": ["revenue"]},
                row_identifier="Konto",
            )
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        "brain.complete",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError(untrusted_response)
        ),
        raising=False,
    )
    result = json.loads(
        ControllingReportTool(tmp_path).safe_run('{"source":"data/actuals.csv"}')
    )

    assert result["mode"] == "deterministic_fallback"
    assert result["summary"] == (
        "Der Bericht basiert auf der geprüften Quelle. "
        "EBITDA im Monat: 10. EGT im Monat: 10."
    )


def _private_narrative_report():
    return {
        "source": "data/KUNDE-ALPHA-PRIVATE.csv",
        "method": "mapping-defined account groups",
        "customer_id": "KUNDE-ALPHA-PRIVATE",
        "file_hash": "sha256:PRIVATE-FILE-HASH",
        "metrics": {
            "ebitda": {
                "month": "10",
                "prior": "9",
                "ytd": "29",
                "account_ids": ["4000"],
            },
            "egt": {
                "month": "8",
                "prior": "7",
                "ytd": "23",
                "account_ids": ["4000", "5000"],
            },
            "KUNDE-ALPHA-PRIVATE-margin": {
                "month": "0.42",
                "account_ids": ["5000"],
            },
        },
        "groups": {
            "revenue": {
                "month": "10",
                "account_ids": ["4000"],
            }
        },
        "material_deviations": [
            {"account_id": "5000", "month_vs_prior": "2"}
        ],
        "reconciliation": {
            "source_account_count": 2,
            "mapped_account_count": 2,
            "ignored_account_count": 0,
            "mapped_source_period_totals": {
                "month": "10",
                "prior": "9",
                "ytd": "29",
            },
            "grouped_period_totals": {
                "month": "10",
                "prior": "9",
                "ytd": "29",
            },
            "customer_id": "KUNDE-ALPHA-PRIVATE",
        },
        "raw_rows": [{"account_id": "4000", "name": "Geheimer Kunde"}],
    }


def _stubbed_narrative_tool(tmp_path, monkeypatch, complete):
    (tmp_path / "project").mkdir()
    (tmp_path / "project" / "manifest.json").write_text(
        '{"model":{"narrative":"cheap"}}', encoding="utf-8"
    )
    report = _private_narrative_report()
    tool = ControllingReportTool(tmp_path)
    monkeypatch.setattr(
        tool._compute, "safe_run", lambda _input: json.dumps(report)
    )
    return tool, report


def test_narrative_ignores_configured_model_response(
    tmp_path, monkeypatch
):
    seen = {}

    def complete(system, user, **_kwargs):
        seen.update(system=system, user=user)
        return '{"summary":"A. B. C.","alarms":[],"actions":["a","b","c"]}'

    tool, _report = _stubbed_narrative_tool(tmp_path, monkeypatch, complete)

    result = json.loads(tool.safe_run('{"source":"data/anything.csv"}'))

    assert seen == {}
    assert result["mode"] == "deterministic_fallback"
    assert result["citations"] == {
        "source": "data/KUNDE-ALPHA-PRIVATE.csv",
        "method": "mapping-defined account groups",
    }


@pytest.mark.parametrize(
    "response",
    [
        '{"summary":"EBITDA ist 999. Lage stabil. Pr\u00fcfung folgt.",'
        '"alarms":[],"actions":["Pr\u00fcfen.","Priorisieren.","Freigeben."]}',
        '{"summary":"Konto 4000 ist auff\u00e4llig. Lage stabil. Pr\u00fcfung folgt.",'
        '"alarms":[],"actions":["Pr\u00fcfen.","Priorisieren.","Freigeben."]}',
        '{"summary":"A. B. C.","alarms":[],'
        '"actions":["Bericht extern versenden.","Priorisieren.","Freigeben."]}',
    ],
    ids=("unsupported-number", "account-id", "external-send"),
)
def test_narrative_rejects_ungrounded_or_external_model_content(
    tmp_path, monkeypatch, response
):
    tool, report = _stubbed_narrative_tool(
        tmp_path,
        monkeypatch,
        lambda _system, _user, **_kwargs: response,
    )

    result = json.loads(tool.safe_run('{"source":"data/anything.csv"}'))

    assert result == ControllingReportTool._fallback(report)
    assert len(result["alarms"]) <= 4
    assert len(result["actions"]) == 3


def test_narrative_rejects_numbers_even_when_present_at_wrong_coordinates(
    tmp_path, monkeypatch
):
    response = json.dumps(
        {
            "summary": (
                "Das EBITDA beträgt 8. Der YTD-Wert beträgt 10. "
                "Die Lage ist stabil."
            ),
            "alarms": [],
            "actions": ["Prüfen.", "Priorisieren.", "Intern freigeben."],
        }
    )
    tool, report = _stubbed_narrative_tool(
        tmp_path,
        monkeypatch,
        lambda _system, _user, **_kwargs: response,
    )

    result = json.loads(tool.safe_run('{"source":"data/anything.csv"}'))

    assert result == ControllingReportTool._fallback(report)


def test_narrative_rejects_numeric_free_model_content_as_untrusted(tmp_path, monkeypatch):
    response = json.dumps(
        {
            "summary": (
                "Die Ergebnislage bleibt beobachtungsbedürftig. "
                "Die Entwicklung wird intern nachverfolgt. "
                "Der deterministische Bericht enthält alle Kennzahlen."
            ),
            "alarms": ["Die Ergebnisentwicklung bleibt im Fokus."],
            "actions": ["Pr\u00fcfen.", "Priorisieren.", "Intern freigeben."],
        }
    )
    tool, report = _stubbed_narrative_tool(
        tmp_path,
        monkeypatch,
        lambda _system, _user, **_kwargs: response,
    )

    result = json.loads(tool.safe_run('{"source":"data/anything.csv"}'))

    assert result == ControllingReportTool._fallback(report)


def _write_strict_actuals(tmp_path, rows):
    (tmp_path / "data").mkdir(exist_ok=True)
    (tmp_path / "data" / "actuals.csv").write_text(
        "Zeilen-ID;Konto;Bezeichnung;Monat;Vormonat;YTD;YTD Vorjahr\n"
        + "\n".join(rows)
        + "\n",
        encoding="utf-8",
    )


@pytest.mark.parametrize(
    ("mutation", "label"),
    [
        (lambda m: m.update(schema_version=2), "schema_version"),
        (lambda m: m["report_metadata"].update(period_start="2026-6-1"), "period_start"),
        (lambda m: m["report_metadata"].update(period_end="2026-05-31"), "period_end"),
        (lambda m: m["report_metadata"].update(currency="Euro"), "currency"),
        (lambda m: m["report_metadata"].update(sign_convention="guess"), "sign_convention"),
        (lambda m: m["report_metadata"].update(row_identifier=""), "row_identifier"),
    ],
)
def test_kmu_report_rejects_invalid_pilot_contract_metadata(
    tmp_path, mutation, label
):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mutation(mapping)
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert label in result


def test_kmu_report_returns_validated_report_metadata_locally(tmp_path):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _write_strict_mapping(tmp_path)

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/actuals.csv"}'
        )
    )

    assert report["report_metadata"] == mapping["report_metadata"]


def test_kmu_report_publishes_mapping_bound_period_labels_and_sign_note(tmp_path):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["columns"].update(
        month="Juni <Ist>", prior="Mai & Vergleich", ytd="2026 kumuliert", ytd_prior="2025 kumuliert"
    )
    _write_strict_mapping(tmp_path, **mapping)
    (tmp_path / "data" / "actuals.csv").write_text(
        "Zeilen-ID;Konto;Bezeichnung;Juni <Ist>;Mai & Vergleich;2026 kumuliert;2025 kumuliert\n"
        "r1;4000;Umsatz;10;9;10;9\n",
        encoding="utf-8",
    )

    report = json.loads(ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    ))

    assert report["period_labels"] == {
        period: mapping["columns"][period]
        for period in ("month", "prior", "ytd", "ytd_prior")
    }
    assert report["sign_convention"] == "source_value_times_group_multiplier"
    assert report["sign_convention_note"].startswith("Vorzeichen:")


@pytest.mark.parametrize(
    "label",
    [" Monat", "Monat ", "Monat\nX", "Monat\u0085X", "Monat\u202eX", "Monat\u200bX", "x" * 81],
)
def test_kmu_report_rejects_unsafe_period_labels(tmp_path, label):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["columns"]["month"] = label
    _write_strict_mapping(tmp_path, **mapping)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")


def test_kmu_report_rejects_casefold_duplicate_period_labels(tmp_path):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["columns"]["prior"] = "monat"
    _write_strict_mapping(tmp_path, **mapping)

    assert ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    ).startswith("FEHLER")


@pytest.mark.parametrize("equivalent", ["Cafe\u0301", "Ｍｏｎａｔ"])
def test_kmu_report_rejects_nfkc_equivalent_column_labels(tmp_path, equivalent):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    mapping = _strict_mapping()
    mapping["columns"]["month"] = "Café" if equivalent == "Cafe\u0301" else "Monat"
    mapping["columns"]["prior"] = equivalent
    _write_strict_mapping(tmp_path, **mapping)

    assert ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    ).startswith("FEHLER")


@pytest.mark.parametrize("label", [" Monat", "Monat ", "Monat\u0085X", "x" * 81])
def test_mapping_schema_matches_runtime_header_bounds(tmp_path, label):
    schema = json.loads(
        (Path(__file__).parents[1] / "presets" / "kmu-controlling" / "pilot" / "account-mapping.schema.json").read_text(encoding="utf-8")
    )
    mapping = _strict_mapping()
    mapping["columns"]["month"] = label
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    _write_strict_mapping(tmp_path, **mapping)

    assert list(Draft202012Validator(schema).iter_errors(mapping))
    assert ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    ).startswith("FEHLER")


def test_kmu_report_rejects_duplicate_row_identifiers(tmp_path):
    _write_strict_actuals(
        tmp_path,
        [
            "duplicate;4000;Umsatz;10;9;10;9",
            "duplicate;4100;Umsatz Export;20;18;20;18",
        ],
    )
    _write_strict_mapping(tmp_path)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "doppelte Zeilen-ID" in result


@pytest.mark.parametrize("missing_index", [3, 4], ids=("month", "prior"))
def test_kmu_report_requires_month_and_prior_for_every_numeric_account(
    tmp_path, missing_index
):
    values = ["r1", "4000", "Umsatz", "10", "9", "10", "9"]
    values[missing_index] = ""
    _write_strict_actuals(tmp_path, [";".join(values)])
    _write_strict_mapping(tmp_path)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "Monat und Vormonat" in result


@pytest.mark.parametrize("missing_index", [5, 6], ids=("ytd", "ytd-prior"))
def test_kmu_report_requires_ytd_periods_for_p_and_l_groups(tmp_path, missing_index):
    values = ["r1", "4000", "Umsatz", "10", "9", "10", "9"]
    values[missing_index] = ""
    _write_strict_actuals(tmp_path, [";".join(values)])
    _write_strict_mapping(tmp_path)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "YTD" in result


def test_kmu_report_allows_blank_ytd_periods_for_balance_groups(tmp_path):
    _write_strict_actuals(
        tmp_path,
        [
            "p1;4000;Umsatz;10;9;10;9",
            "b1;1000;Bank;20;19;;",
        ],
    )
    _write_strict_mapping(tmp_path)

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/actuals.csv"}'
        )
    )

    assert report["metrics"]["liquidity"]["ytd"] is None
    assert report["metrics"]["liquidity"]["ytd_prior"] is None


def test_balance_group_propagates_partial_ytd_to_metric_and_reconciliation(tmp_path):
    _write_strict_actuals(
        tmp_path,
        [
            "p1;4000;Umsatz;10;9;30;27",
            "b1;1000;Bank A;20;19;60;57",
            "b2;1001;Bank B;25;24;;",
        ],
    )
    mapping = _write_strict_mapping(tmp_path)

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/actuals.csv"}'
        )
    )

    assert report["groups"]["liquidity"]["ytd"] is None
    assert report["groups"]["liquidity"]["ytd_prior"] is None
    assert report["metrics"]["liquidity"]["ytd"] is None
    assert report["metrics"]["liquidity"]["ytd_prior"] is None
    assert report["reconciliation"]["mapped_source_period_totals"]["ytd"] is None
    assert report["reconciliation"]["grouped_period_totals"]["ytd"] is None
    assert artifacts_module._sourced_totals_pass(report, mapping)


def test_free_metric_propagates_missing_balance_group_ytd_to_acceptance(tmp_path):
    groups = {
        "revenue": {"prefixes": ["4"], "multiplier": 1},
        "liquidity": {"prefixes": ["1"], "multiplier": 1},
    }
    metrics = {
        "ebitda": ["revenue"],
        "depreciation": ["revenue"],
        "egt": ["revenue"],
        "free_kpi": ["revenue", "liquidity"],
    }
    _write_strict_actuals(
        tmp_path,
        [
            "p1;4000;Umsatz;10;9;30;27",
            "b1;1000;Bank;20;19;;",
        ],
    )
    mapping = _write_strict_mapping(tmp_path, groups=groups, metrics=metrics)

    report = json.loads(
        ControllingComputeTool(tmp_path).safe_run(
            '{"operation":"kmu_report","source":"data/actuals.csv"}'
        )
    )

    assert report["metrics"]["free_kpi"]["month"] == "30"
    assert report["metrics"]["free_kpi"]["ytd"] is None
    assert report["metrics"]["free_kpi"]["ytd_prior"] is None
    assert artifacts_module._sourced_totals_pass(report, mapping)

    partial_ytd = json.loads(json.dumps(report))
    partial_ytd["metrics"]["free_kpi"]["ytd"] = "30"
    assert not artifacts_module._sourced_totals_pass(partial_ytd, mapping)


def test_kmu_report_rejects_mapping_symlink_outside_instance(tmp_path, monkeypatch):
    _write_strict_actuals(tmp_path, ["r1;4000;Umsatz;10;9;10;9"])
    pilot = tmp_path / "project" / "pilot"
    pilot.mkdir(parents=True)
    outside = tmp_path.parent / f"{tmp_path.name}-outside-mapping.json"
    outside.write_text(json.dumps(_strict_mapping()), encoding="utf-8")
    declared = pilot / "account-mapping.json"
    tool = ControllingComputeTool(tmp_path)
    try:
        declared.symlink_to(outside)
    except OSError:
        original_resolve = type(tmp_path).resolve

        def resolve(path, *, strict=False):
            if path == declared:
                return outside
            return original_resolve(path, strict=strict)

        monkeypatch.setattr(type(tmp_path), "resolve", resolve)

    result = tool.safe_run(
        '{"operation":"kmu_report","source":"data/actuals.csv"}'
    )

    assert result.startswith("FEHLER")
    assert "Instanz" in result


def test_malformed_xlsx_returns_fail_closed_error(tmp_path):
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "broken.xlsx").write_bytes(b"not-a-zip")

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/broken.xlsx","column":"value"}'
    )

    assert result.startswith("FEHLER: XLSX")


def test_xlsx_member_count_limit_is_enforced_before_parsing(tmp_path, monkeypatch):
    (tmp_path / "data").mkdir()
    source = tmp_path / "data" / "many-members.xlsx"
    with ZipFile(source, "w", compression=ZIP_DEFLATED) as archive:
        for index in range(3):
            archive.writestr(f"member-{index}.xml", "x")
    monkeypatch.setattr(controlling_tools, "_MAX_XLSX_MEMBERS", 2, raising=False)

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/many-members.xlsx","column":"value"}'
    )

    assert result.startswith("FEHLER: XLSX enthält zu viele Archivdateien")


def test_xlsx_expansion_limit_is_enforced_before_parsing(tmp_path, monkeypatch):
    (tmp_path / "data").mkdir()
    source = tmp_path / "data" / "expanded.xlsx"
    with ZipFile(source, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", "x" * 100)
    monkeypatch.setattr(
        controlling_tools, "_MAX_XLSX_UNCOMPRESSED_BYTES", 50, raising=False
    )

    result = ControllingComputeTool(tmp_path).safe_run(
        '{"operation":"sum","source":"data/expanded.xlsx","column":"value"}'
    )

    assert result.startswith("FEHLER: XLSX entpackt zu groß")
