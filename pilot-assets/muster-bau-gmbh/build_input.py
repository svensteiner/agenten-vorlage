#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Baut die MUSTER-EINGABEDATEI fuer den Pilot-Probelauf "Muster Bau GmbH".

Erzeugt ein mehrblaettriges Controlling-Excel, wie es eine mittelstaendische
Bau-GmbH fuehren wuerde — MIT der Gelb-Konvention (FFF2CC) auf den echten
Eingabezellen und MIT blattuebergreifenden Formeln (=Annahmen!B5). Das ist die
"Kunden-Dateistruktur rein" aus SYNTHETIC_DATA_PIPELINE.md, Schritt 1.

Dieses Skript ersetzt KEIN echtes Kundendokument — es baut eine rein fiktive
Struktur, aus der der synthetic-data-creator (Schritt 2) den synthetischen
Zwilling erzeugt. Reproduzierbar: gleicher Aufbau bei jedem Lauf.

    python build_input.py            # schreibt muster-bau-controlling_vorlage.xlsx

Bewusst gewaehlt (Lernpunkt fuer echte Pilots):
- Echte Eingaben (Mengen, Einzelpreise, %-Annahmen) sind GELB -> werden
  synthetisiert.
- Gesetzliche Konstanten (USt-Satz 19 %) sind NICHT gelb -> bleiben unangetastet.
- Alle Summen/Aufschlaege sind FORMELN -> das Modell rechnet nach dem Klonen
  unveraendert weiter.
"""
from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

INPUT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "muster-bau-controlling_vorlage.xlsx")


def _input(ws, coord, value):
    """Setzt einen Eingabewert und markiert die Zelle gelb (Gelb-Konvention)."""
    ws[coord] = value
    ws[coord].fill = INPUT_FILL


def build() -> str:
    wb = Workbook()

    # --- Blatt 1: Annahmen ------------------------------------------------
    an = wb.active
    an.title = "Annahmen"
    an["A1"] = "Muster Bau GmbH - Controlling-Annahmen 2026"
    an["A1"].font = TITLE
    an["A3"] = "Parameter"
    an["A3"].font = BOLD
    an["B3"] = "Wert"
    an["B3"].font = BOLD
    an["A4"] = "USt-Satz (%)"
    an["B4"] = 19                       # gesetzliche Konstante -> NICHT gelb
    an["A5"] = "Gemeinkostenumlage (%)"
    _input(an, "B5", 8.5)               # echte Annahme -> gelb
    an["A6"] = "Materialaufschlag (%)"
    _input(an, "B6", 12)                # echte Annahme -> gelb
    an.column_dimensions["A"].width = 28
    an.column_dimensions["B"].width = 12

    # --- Blatt 2: Monatsabschluss Mai ------------------------------------
    mo = wb.create_sheet("Monatsabschluss Mai")
    mo["A1"] = "Muster Bau GmbH - Monatsabschluss Mai 2026 (BAB Kostenstelle 100)"
    mo["A1"].font = TITLE
    for col, head in zip("ABCDE", ["Kostenart", "Menge", "Einheit", "Einzelpreis (EUR)", "Summe (EUR)"]):
        mo[f"{col}3"] = head
        mo[f"{col}3"].font = BOLD

    # Datenzeilen: Menge (B) und Einzelpreis (D) sind gelbe Eingaben,
    # Summe (E) ist eine Formel = Menge * Einzelpreis.
    rows = [
        ("Rohbau Beton/Stahl",     420,  "m3",  190),
        ("Bewehrungsstahl",         38,  "t",   1180),
        ("Zimmererarbeiten",       620,  "h",   58),
        ("Elektroinstallation",    340,  "h",   62),
        ("Sanitaer/Heizung",       280,  "h",   66),
        ("Baustelleneinrichtung",    1,  "pausch", 48000),
    ]
    r = 4
    for name, menge, einheit, preis in rows:
        mo[f"A{r}"] = name
        _input(mo, f"B{r}", menge)
        mo[f"C{r}"] = einheit
        _input(mo, f"D{r}", preis)
        mo[f"E{r}"] = f"=B{r}*D{r}"     # Formel bleibt beim Klonen erhalten
        r += 1

    last = r - 1                        # letzte Datenzeile (=9)
    r += 1                              # eine Leerzeile
    mo[f"A{r}"] = "Summe Direktkosten netto"
    mo[f"A{r}"].font = BOLD
    mo[f"E{r}"] = f"=SUM(E4:E{last})"
    direkt = r
    r += 1
    mo[f"A{r}"] = "Gemeinkostenumlage"
    mo[f"E{r}"] = f"=E{direkt}*Annahmen!B5/100"   # blattuebergreifend
    gk = r
    r += 1
    mo[f"A{r}"] = "Materialaufschlag"
    mo[f"E{r}"] = f"=E{direkt}*Annahmen!B6/100"   # blattuebergreifend
    ma = r
    r += 1
    mo[f"A{r}"] = "Netto gesamt"
    mo[f"A{r}"].font = BOLD
    mo[f"E{r}"] = f"=E{direkt}+E{gk}+E{ma}"
    netto = r
    r += 1
    mo[f"A{r}"] = "USt"
    mo[f"E{r}"] = f"=E{netto}*Annahmen!B4/100"    # blattuebergreifend auf Konstante
    ust = r
    r += 1
    mo[f"A{r}"] = "Bruttosumme"
    mo[f"A{r}"].font = BOLD
    mo[f"E{r}"] = f"=E{netto}+E{ust}"

    mo.column_dimensions["A"].width = 30
    for c in "BCDE":
        mo.column_dimensions[c].width = 16

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Muster-Eingabedatei geschrieben: {path}")
