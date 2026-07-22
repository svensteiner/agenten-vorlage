"""Fail-closed, source-backed capabilities for a controlling instance."""

import csv
from datetime import date
from decimal import Decimal, InvalidOperation
import hashlib
from itertools import islice
import json
from pathlib import Path
import re
from typing import Any
import unicodedata
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, LargeZipFile, ZipFile

_INPUT_DIRS = ("inbox", "data", "imports")
_SUPPORTED_SUFFIXES = {".csv", ".xlsx"}
_PERIODS = ("month", "prior", "ytd")
_MAX_REQUEST_CHARS = 64 * 1024
_MAX_MAPPING_BYTES = 1024 * 1024
_MAX_SOURCE_BYTES = 25 * 1024 * 1024
_MAX_SOURCE_ROWS = 100_000
_MAX_SOURCE_COLUMNS = 256
_MAX_SOURCE_GLOBS = 20
_MAX_SOURCE_GLOB_CHARS = 512
_MAX_XLSX_MEMBERS = 2_048
_MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_XLSX_MEMBER_BYTES = 32 * 1024 * 1024
_SIGN_CONVENTION = "source_value_times_group_multiplier"
_SIGN_FORMULA = "normalized_value = source_value * group.multiplier"
_SIGN_CONVENTION_NOTE = (
    "Vorzeichen: Quellwert × Gruppen-Multiplikator; positive Werte erhöhen "
    "die Kennzahl, negative Werte verringern sie."
)
_CALCULATION_VERSION = "kmu-report-v1"
PNL_METRIC_NAMES = frozenset(
    {
        "ebitda",
        "ebit",
        "egt",
        "ebt",
        "depreciation",
        "operating_result",
        "net_income",
    }
)
REPORT_PERIODS = ("month", "prior", "ytd", "ytd_prior")
MANDATORY_MAPPING_METRICS = frozenset({"ebitda", "depreciation", "egt"})
DERIVED_RATIO_METRICS = (
    ("material_ratio", "material", "revenue"),
    ("personnel_ratio", "personnel", "revenue"),
)
DIRECT_BALANCE_METRICS = (
    ("receivables", "receivables"),
    ("liquidity", "liquidity"),
    ("payables", "payables"),
)
RESERVED_DERIVED_METRIC_NAMES = frozenset(
    metric_name
    for metric_name, *_groups in (*DERIVED_RATIO_METRICS, *DIRECT_BALANCE_METRICS)
)


def is_safe_report_label(value: Any) -> bool:
    """Return whether a source header is safe to publish in a report view."""
    return (
        isinstance(value, str)
        and value == value.strip()
        and 0 < len(value) <= 80
        and not any(unicodedata.category(char).startswith("C") for char in value)
    )


def labels_are_unique(values: list[Any]) -> bool:
    """Compare report labels using the same Unicode identity seen by readers."""
    return len(
        {
            unicodedata.normalize("NFKC", value).casefold()
            for value in values
            if isinstance(value, str)
        }
    ) == len(values)


class _InstanceInputCapability:
    def __init__(self, instance_root: Path):
        self._root = instance_root.resolve()

    def _resolve_source(self, source: str) -> tuple[Path | None, str | None]:
        if not isinstance(source, str) or not source.strip():
            return None, "Quelle fehlt."
        path = (self._root / source).resolve()
        try:
            relative = path.relative_to(self._root)
        except ValueError:
            return None, "Quelle liegt außerhalb der Instanz."
        if not relative.parts or relative.parts[0] not in _INPUT_DIRS:
            return None, "Quelle liegt nicht in einem freigegebenen Eingabeordner."
        if path.suffix.lower() not in _SUPPORTED_SUFFIXES:
            return None, "Nur CSV oder XLSX sind erlaubt."
        try:
            is_file = path.is_file()
            size = path.stat().st_size if is_file else 0
        except OSError:
            return None, "Quelle existiert nicht oder ist keine Datei."
        if not is_file:
            return None, "Quelle existiert nicht oder ist keine Datei."
        if size > _MAX_SOURCE_BYTES:
            return None, "Eingabedatei ist zu groß."
        return path, None

    def _source_label(self, path: Path) -> str:
        return path.relative_to(self._root).as_posix()


class InstanceReadTool(_InstanceInputCapability):
    """Exposes source metadata only; ledger values never enter the model context."""

    description = (
        "Lists approved inputs and returns schema/sheet/row-count metadata only."
    )

    def safe_run(self, input_str: str) -> str:
        command = (input_str or "").strip()
        if command == "list":
            found = [
                self._source_label(p)
                for directory in _INPUT_DIRS
                for p in (self._root / directory).rglob("*")
                if (self._root / directory).is_dir()
                and p.is_file()
                and p.suffix.lower() in _SUPPORTED_SUFFIXES
            ]
            return (
                "SOURCES:\n" + "\n".join(sorted(found)) if found else "SOURCES: (none)"
            )
        if not command.startswith("read:"):
            return "FEHLER: Erlaubt sind 'list' oder 'read:<quelle>'."
        path, error = self._resolve_source(command[5:].strip())
        if error:
            return f"FEHLER: {error}"
        try:
            rows, columns, sheets = ControllingComputeTool._metadata(path)
        except ValueError as exc:
            return f"FEHLER: {exc}"
        return "\n".join(
            (
                f"SOURCE: {self._source_label(path)}",
                f"FORMAT: {path.suffix[1:].upper()}",
                f"SHEETS: {', '.join(sheets)}",
                f"ROWS: {rows}",
                f"SCHEMA: {', '.join(columns)}",
            )
        )


class ControllingComputeTool(_InstanceInputCapability):
    description = "Deterministic, source-backed sums, comparisons, and mapping-defined KMU reports."

    def safe_run(self, input_str: str) -> str:
        if not isinstance(input_str, str) or len(input_str) > _MAX_REQUEST_CHARS:
            return "FEHLER: Berechnung muss ein begrenztes JSON-Objekt sein."
        try:
            request = json.loads(input_str)
        except (TypeError, json.JSONDecodeError, RecursionError):
            return "FEHLER: Berechnung muss ein JSON-Objekt sein."
        if not isinstance(request, dict):
            return "FEHLER: Berechnung muss ein JSON-Objekt sein."
        operation, source = request.get("operation"), request.get("source", "")
        mapping = None
        try:
            if operation == "kmu_report":
                mapping = self._mapping()
            path, error = (
                self._single_source(mapping)
                if operation == "kmu_report" and source == "auto"
                else self._resolve_source(source)
            )
        except ValueError as exc:
            return f"FEHLER: {exc}"
        if error:
            return f"FEHLER: {error}"
        try:
            if operation == "kmu_report":
                return self._kmu_report(path, mapping)
            rows = self._rows(path)
            if operation == "sum":
                column = self._column(request, "column")
                result = self._sum(rows, column)
                method = f"sum({column})"
            elif operation == "difference":
                actual, budget = (
                    self._column(request, "actual_column"),
                    self._column(request, "budget_column"),
                )
                result = self._sum(rows, actual) - self._sum(rows, budget)
                method = f"sum({actual}) - sum({budget})"
            else:
                return "FEHLER: Erlaubte Berechnungen sind 'sum', 'difference' und 'kmu_report'."
        except ValueError as exc:
            return f"FEHLER: {exc}"
        return f"RESULT: {self._format(result)}\nSOURCE: {self._source_label(path)}\nMETHOD: {method}\nROWS: {len(rows)}"

    def _single_source(
        self, mapping: dict[str, Any] | None = None
    ) -> tuple[Path | None, str | None]:
        globs = (mapping or {}).get("source_globs")
        try:
            candidates = (
                [
                    path
                    for pattern in globs
                    for path in self._root.glob(pattern)
                    if path.suffix.lower() in _SUPPORTED_SUFFIXES
                ]
                if globs
                else [
                    path
                    for directory in _INPUT_DIRS
                    if (self._root / directory).is_dir()
                    for path in (self._root / directory).rglob("*")
                    if path.suffix.lower() in _SUPPORTED_SUFFIXES
                ]
            )
        except (OSError, ValueError):
            return None, "source_globs enthält ein ungültiges Suchmuster."
        paths = []
        for candidate in candidates:
            try:
                source = candidate.relative_to(self._root).as_posix()
            except ValueError:
                return None, "Automatische Quelle liegt außerhalb der Instanz."
            path, error = self._resolve_source(source)
            if error:
                return None, error
            if path not in paths:
                paths.append(path)
        return (
            (paths[0], None)
            if len(paths) == 1
            else (None, "Automatische Quelle erfordert genau eine CSV/XLSX-Eingabe.")
        )

    def _kmu_report(
        self, path: Path, mapping: dict[str, Any] | None = None
    ) -> str:
        mapping = mapping or self._mapping()
        source_fingerprint = self._file_sha256(path, max_bytes=_MAX_SOURCE_BYTES)
        required_headers = list(
            dict.fromkeys(
                [
                    *mapping["columns"].values(),
                    mapping["report_metadata"]["row_identifier"],
                ]
            )
        )
        rows = self._rows(path, required_headers)
        groups, reconciliation = self._groups(rows, mapping)
        metrics = {
            name: self._metric(groups, refs)
            for name, refs in mapping["metrics"].items()
        }
        for metric_name, numerator, denominator in DERIVED_RATIO_METRICS:
            if numerator in groups:
                ratio_periods = tuple(
                    period for period in groups[numerator] if period != "accounts"
                )
                metrics[metric_name] = self._ratio(
                    groups[numerator],
                    groups.get(denominator, self._empty(ratio_periods)),
                )
        for metric_name, group in DIRECT_BALANCE_METRICS:
            if group in groups:
                metrics[metric_name] = groups[group]
        report = {
            "source": self._source_label(path),
            "method": "mapping-defined account groups; month/prior/ytd aggregation",
            "report_metadata": mapping["report_metadata"],
            "period_labels": {
                period: mapping["columns"][period] for period in REPORT_PERIODS
            },
            "sign_convention": _SIGN_CONVENTION,
            "sign_convention_note": _SIGN_CONVENTION_NOTE,
            "metrics": {k: self._public(v) for k, v in metrics.items()},
            "groups": {
                k: self._public(v, include_row_count=True) for k, v in groups.items()
            },
            "material_deviations": self._deviations(rows, mapping),
            "reconciliation": reconciliation,
            "provenance": {
                "input_fingerprint": source_fingerprint,
                "mapping_version": mapping["_mapping_version"],
                "calculation_version": _CALCULATION_VERSION,
            },
        }
        if self._file_sha256(path, max_bytes=_MAX_SOURCE_BYTES) != source_fingerprint:
            raise ValueError("Eingabedatei wurde während der Berechnung verändert.")
        if (
            self._file_sha256(
                mapping["_mapping_path"], max_bytes=_MAX_MAPPING_BYTES
            )
            != mapping["_mapping_version"]
        ):
            raise ValueError("Konten-Mapping wurde während der Berechnung verändert.")
        return json.dumps(
            report,
            ensure_ascii=False,
            separators=(",", ":"),
        )

    def _mapping(self) -> dict[str, Any]:
        declared_mapping_path = (
            self._root / "project" / "pilot" / "account-mapping.json"
        )
        try:
            mapping_path = declared_mapping_path.resolve(strict=True)
            mapping_path.relative_to(self._root)
            mapping_size = mapping_path.stat().st_size
        except ValueError:
            raise ValueError("Konten-Mapping liegt außerhalb der Instanz.") from None
        except OSError:
            raise ValueError(
                "Gültiges project/pilot/account-mapping.json fehlt."
            ) from None
        if mapping_size > _MAX_MAPPING_BYTES:
            raise ValueError("Konten-Mapping ist zu groß.")
        try:
            mapping_bytes = mapping_path.read_bytes()
            if len(mapping_bytes) > _MAX_MAPPING_BYTES:
                raise ValueError("Konten-Mapping ist zu groß.")
            mapping_text = mapping_bytes.decode("utf-8")
        except (OSError, UnicodeError):
            raise ValueError(
                "Gültiges project/pilot/account-mapping.json fehlt."
            ) from None
        try:
            raw = json.loads(mapping_text)
        except (json.JSONDecodeError, RecursionError):
            raise ValueError(
                "Gültiges project/pilot/account-mapping.json fehlt."
            ) from None
        if not isinstance(raw, dict):
            raise ValueError("Konten-Mapping muss ein JSON-Objekt sein.")
        expected_top_level = {
            "schema_version",
            "report_metadata",
            "source_globs",
            "columns",
            "sign_semantics",
            "groups",
            "metrics",
            "thresholds",
        }
        missing_top_level = expected_top_level - set(raw)
        if missing_top_level:
            raise ValueError(
                "Konten-Mapping fehlt: " + ", ".join(sorted(missing_top_level))
            )
        if set(raw) != expected_top_level:
            raise ValueError("Konten-Mapping enthält unbekannte Top-Level-Felder.")
        report_metadata = self._validated_report_metadata(raw)
        self._validate_sign_semantics(raw.get("sign_semantics"))
        columns, groups, metrics, thresholds = (
            raw.get(k) for k in ("columns", "groups", "metrics", "thresholds")
        )
        required = {"account_id", "account_name", *_PERIODS, "ytd_prior"}
        if not isinstance(columns, dict) or set(columns) != required:
            raise ValueError("Konten-Mapping hat keine vollständige Spaltenzuordnung.")
        if not self._valid_column_headers(columns):
            raise ValueError(
                "Jede semantische Spalte muss eindeutig zugeordnet sein."
            )
        if (
            not isinstance(groups, dict)
            or not groups
            or not isinstance(metrics, dict)
            or not metrics
            or not isinstance(thresholds, dict)
        ):
            raise ValueError(
                "Konten-Mapping benötigt Gruppen, Kennzahlen und Schwellenwerte."
            )
        for rule in groups.values():
            if not isinstance(rule, dict):
                raise ValueError(
                    "Konten-Mapping enthält eine ungültige Gruppenregel."
                )
            matchers = []
            for key in ("prefixes", "name_matches"):
                if key not in rule:
                    continue
                values = rule[key]
                if (
                    not isinstance(values, list)
                    or not values
                    or not all(
                        isinstance(value, str) and value.strip() for value in values
                    )
                ):
                    raise ValueError(
                        "Konten-Mapping enthält eine ungültige Gruppenregel."
                    )
                matchers.append(key)
            if set(rule) - {"prefixes", "name_matches", "multiplier", "ignored"}:
                raise ValueError(
                    "Konten-Mapping enthält eine ungültige Gruppenregel."
                )
            if "multiplier" not in rule:
                raise ValueError("Jede Gruppenregel benötigt einen multiplier.")
            if not matchers or (
                "ignored" in rule and not isinstance(rule["ignored"], bool)
            ):
                raise ValueError(
                    "Konten-Mapping enthält eine ungültige Gruppenregel."
                )
            multiplier = self._finite_decimal(
                rule["multiplier"], "Gruppen-Multiplikator"
            )
            if multiplier == 0:
                raise ValueError(
                    "multiplier darf nicht 0 sein; Ausschlüsse verwenden "
                    "ignored:true mit einem von 0 verschiedenen multiplier."
                )
        required_metrics = {"ebitda", "depreciation", "egt"}
        if not required_metrics <= set(metrics):
            raise ValueError(
                "Konten-Mapping benötigt die Pflichtkennzahlen ebitda, depreciation und egt."
            )
        if RESERVED_DERIVED_METRIC_NAMES.intersection(metrics):
            raise ValueError(
                "Konten-Mapping verwendet einen reservierten abgeleiteten Kennzahlnamen."
            )
        if not all(
            isinstance(v, list)
            and v
            and len(v) == len(set(v))
            and all(isinstance(r, str) and r in groups for r in v)
            for v in metrics.values()
        ):
            raise ValueError("Konten-Mapping enthält ungültige Kennzahlgruppen.")
        if any(
            groups[ref].get("ignored") is True
            for refs in metrics.values()
            for ref in refs
        ):
            raise ValueError(
                "Kennzahlen dürfen keine bewusst ignorierte Gruppe referenzieren."
            )
        try:
            threshold = self._finite_decimal(
                thresholds["material_deviation_abs"], "material_deviation_abs"
            )
        except KeyError:
            raise ValueError(
                "Konten-Mapping benötigt einen numerischen material_deviation_abs."
            ) from None
        if set(thresholds) != {"material_deviation_abs"} or threshold < 0:
            raise ValueError("material_deviation_abs muss eine nichtnegative Zahl sein.")
        source_globs = raw.get("source_globs")
        if (
            not isinstance(source_globs, list)
            or not source_globs
            or len(source_globs) > _MAX_SOURCE_GLOBS
            or len(source_globs) != len(set(source_globs))
        ):
            raise ValueError("source_globs muss eine eindeutige, nichtleere Liste sein.")
        for pattern in source_globs:
            candidate = Path(pattern) if isinstance(pattern, str) else None
            if (
                candidate is None
                or not pattern
                or len(pattern) > _MAX_SOURCE_GLOB_CHARS
                or "\\" in pattern
                or candidate.is_absolute()
                or ".." in candidate.parts
                or not candidate.parts
                or candidate.parts[0] not in _INPUT_DIRS
            ):
                raise ValueError("source_globs darf die Instanz nicht verlassen.")
        return {
            "schema_version": 1,
            "report_metadata": report_metadata,
            "columns": columns,
            "groups": groups,
            "metrics": metrics,
            "thresholds": thresholds,
            "source_globs": source_globs,
            "_mapping_path": mapping_path,
            "_mapping_version": hashlib.sha256(mapping_bytes).hexdigest(),
        }

    @staticmethod
    def _valid_column_headers(columns: dict[str, Any]) -> bool:
        """Reject ambiguous or unsafe source headers before they become report labels."""
        headers = list(columns.values())
        return (
            all(is_safe_report_label(value) for value in headers)
            and labels_are_unique(headers)
        )

    @staticmethod
    def _file_sha256(path: Path, *, max_bytes: int) -> str:
        digest = hashlib.sha256()
        total = 0
        try:
            with path.open("rb") as handle:
                for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                    total += len(chunk)
                    if total > max_bytes:
                        raise ValueError("Quelldatei wurde während des Lesens zu groß.")
                    digest.update(chunk)
        except OSError:
            raise ValueError("Quelldatei konnte nicht stabil gelesen werden.") from None
        return digest.hexdigest()

    @staticmethod
    def _validated_report_metadata(raw: dict[str, Any]) -> dict[str, str]:
        schema_version = raw.get("schema_version")
        if (
            not isinstance(schema_version, int)
            or isinstance(schema_version, bool)
            or schema_version != 1
        ):
            raise ValueError("schema_version muss 1 sein.")
        metadata = raw.get("report_metadata")
        required = {
            "period_start",
            "period_end",
            "currency",
            "sign_convention",
            "row_identifier",
        }
        if not isinstance(metadata, dict) or set(metadata) != required:
            raise ValueError(
                "report_metadata muss period_start, period_end, currency, "
                "sign_convention und row_identifier enthalten."
            )
        dates: dict[str, date] = {}
        for field in ("period_start", "period_end"):
            value = metadata.get(field)
            if not isinstance(value, str) or not re.fullmatch(
                r"\d{4}-\d{2}-\d{2}", value
            ):
                raise ValueError(f"{field} muss ein ISO-Datum YYYY-MM-DD sein.")
            try:
                dates[field] = date.fromisoformat(value)
            except ValueError:
                raise ValueError(
                    f"{field} muss ein gültiges ISO-Datum YYYY-MM-DD sein."
                ) from None
        if dates["period_end"] < dates["period_start"]:
            raise ValueError("period_end darf nicht vor period_start liegen.")
        currency = metadata.get("currency")
        if not isinstance(currency, str) or not re.fullmatch(r"[A-Z]{3}", currency):
            raise ValueError("currency muss ein dreistelliger Großbuchstabencode sein.")
        if metadata.get("sign_convention") != _SIGN_CONVENTION:
            raise ValueError(
                f"sign_convention muss '{_SIGN_CONVENTION}' sein."
            )
        row_identifier = metadata.get("row_identifier")
        if not isinstance(row_identifier, str) or not row_identifier.strip():
            raise ValueError("row_identifier muss eine nichtleere Spalte benennen.")
        return {field: metadata[field] for field in sorted(required)}

    @staticmethod
    def _validate_sign_semantics(value: Any) -> None:
        required = {"formula", "positive_contribution", "negative_contribution"}
        if not isinstance(value, dict) or set(value) != required:
            raise ValueError("sign_semantics fehlt oder ist ungültig.")
        if value.get("formula") != _SIGN_FORMULA:
            raise ValueError(f"sign_semantics.formula muss '{_SIGN_FORMULA}' sein.")
        for field in ("positive_contribution", "negative_contribution"):
            description = value.get(field)
            if not isinstance(description, str) or not description.strip():
                raise ValueError(f"sign_semantics.{field} muss beschrieben sein.")

    def _groups(
        self, rows: list[dict[str, str]], mapping: dict[str, Any]
    ) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
        cols = mapping["columns"]
        periods = self._periods(mapping)
        out = {
            name: self._empty(periods)
            for name, rule in mapping["groups"].items()
            if rule.get("ignored") is not True
        }
        source_totals = {period: None for period in periods}
        grouped_totals = {period: None for period in periods}
        source_account_count = mapped_account_count = ignored_account_count = 0
        row_identifier = mapping["report_metadata"]["row_identifier"]
        seen_row_ids: set[str] = set()
        pnl_groups = {
            group
            for metric, refs in mapping["metrics"].items()
            if metric.casefold() in PNL_METRIC_NAMES
            for group in refs
        }
        for row in rows:
            account, name = (
                row[cols["account_id"]].strip(),
                row[cols["account_name"]].strip(),
            )
            if not self._is_account_id(account):
                continue
            row_id = row[row_identifier].strip()
            if not row_id:
                raise ValueError(f"Zeilen-ID fehlt für Konto '{account}'.")
            if row_id in seen_row_ids:
                raise ValueError(f"doppelte Zeilen-ID '{row_id}'.")
            seen_row_ids.add(row_id)
            values = {
                period: self._optional_decimal(row[cols[period]], account)
                for period in periods
            }
            group, rule = self._matched_group(account, name, mapping)
            if values["month"] is None or values["prior"] is None:
                raise ValueError(
                    f"Monat und Vormonat fehlen für Konto '{account}'."
                )
            if group in pnl_groups and any(
                values[period] is None for period in ("ytd", "ytd_prior")
            ):
                raise ValueError(
                    f"YTD und YTD-Vorjahr fehlen für GuV-Konto '{account}'."
                )
            source_account_count += 1
            if rule.get("ignored") is True:
                ignored_account_count += 1
                continue
            mapped_account_count += 1
            mult = self._finite_decimal(
                rule.get("multiplier", 1), "Gruppen-Multiplikator"
            )
            for period in periods:
                source_totals[period] = self._add(
                    source_totals[period], values[period], mult
                )
                grouped_totals[period] = self._add(
                    grouped_totals[period], values[period], mult
                )
                out[group][period] = self._add(out[group][period], values[period], mult)
            out[group]["accounts"].append(account)
        return out, {
            "source_account_count": source_account_count,
            "mapped_account_count": mapped_account_count,
            "ignored_account_count": ignored_account_count,
            "mapped_source_period_totals": {
                period: self._format(value) if value is not None else None
                for period, value in source_totals.items()
            },
            "grouped_period_totals": {
                period: self._format(value) if value is not None else None
                for period, value in grouped_totals.items()
            },
        }

    @staticmethod
    def _is_account_id(value: str) -> bool:
        return bool(re.fullmatch(r"\d[\d .-]*", value))

    @staticmethod
    def _matches(account: str, name: str, rule: dict[str, Any]) -> bool:
        return any(account.startswith(str(x)) for x in rule.get("prefixes", [])) or any(
            str(x).casefold() in name.casefold() for x in rule.get("name_matches", [])
        )

    def _matched_group(
        self, account: str, name: str, mapping: dict[str, Any]
    ) -> tuple[str, dict[str, Any]]:
        matches = [
            (group, rule)
            for group, rule in mapping["groups"].items()
            if self._matches(account, name, rule)
        ]
        if not matches:
            raise ValueError(f"Konto '{account}' ist nicht zugeordnet.")
        if len(matches) != 1:
            raise ValueError(f"Konto '{account}' ist mehrfach zugeordnet.")
        return matches[0]

    @staticmethod
    def _empty(periods: tuple[str, ...] = _PERIODS) -> dict[str, Any]:
        return {**{period: None for period in periods}, "accounts": []}

    @staticmethod
    def _periods(mapping: dict[str, Any]) -> tuple[str, ...]:
        return _PERIODS + (("ytd_prior",) if "ytd_prior" in mapping["columns"] else ())

    @staticmethod
    def _add(
        current: Decimal | None, value: Decimal | None, multiplier: Decimal
    ) -> Decimal | None:
        if value is None:
            return current
        return value * multiplier if current is None else current + (value * multiplier)

    def _metric(
        self, groups: dict[str, dict[str, Any]], refs: list[str]
    ) -> dict[str, Any]:
        periods = tuple(period for period in groups[refs[0]] if period != "accounts")
        out = self._empty(periods)
        for ref in refs:
            for period in periods:
                out[period] = self._add(out[period], groups[ref][period], Decimal("1"))
            out["accounts"].extend(groups[ref]["accounts"])
        return out

    @staticmethod
    def _ratio(a: dict[str, Any], b: dict[str, Any]) -> dict[str, Any]:
        return {
            **{
                period: (
                    abs(a[period]) / abs(b[period])
                    if a[period] is not None and b[period]
                    else None
                )
                for period in a
                if period != "accounts"
            },
            "accounts": a["accounts"] + b["accounts"],
        }

    def _deviations(
        self, rows: list[dict[str, str]], mapping: dict[str, Any]
    ) -> list[dict[str, str]]:
        cols = mapping["columns"]
        threshold = abs(
            self._finite_decimal(
                mapping["thresholds"]["material_deviation_abs"],
                "material_deviation_abs",
            )
        )
        out = []
        for row in rows:
            account = row[cols["account_id"]].strip()
            if not self._is_account_id(account):
                continue
            _, rule = self._matched_group(
                account, row[cols["account_name"]].strip(), mapping
            )
            if rule.get("ignored") is True:
                continue
            month = self._optional_decimal(row[cols["month"]], account)
            prior = self._optional_decimal(row[cols["prior"]], account)
            if month is None or prior is None:
                continue
            mult = self._finite_decimal(
                rule.get("multiplier", 1), "Gruppen-Multiplikator"
            )
            delta = (month - prior) * mult
            if abs(delta) >= threshold:
                out.append(
                    {"account_id": account, "month_vs_prior": self._format(delta)}
                )
        return sorted(
            out,
            key=lambda x: (abs(Decimal(x["month_vs_prior"])), x["account_id"]),
            reverse=True,
        )

    def _public(
        self, value: dict[str, Any], *, include_row_count: bool = False
    ) -> dict[str, Any]:
        public = {
            **{
                p: (self._format(value[p]) if value[p] is not None else None)
                for p in value
                if p != "accounts"
            },
            "account_ids": sorted(set(value["accounts"])),
        }
        if include_row_count:
            public["row_count"] = len(value["accounts"])
        return public

    @staticmethod
    def _column(request: dict[str, Any], key: str) -> str:
        value = request.get(key)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"Spalte '{key}' fehlt.")
        return value.strip()

    @classmethod
    def _sum(cls, rows: list[dict[str, str]], column: str) -> Decimal:
        if not rows or column not in rows[0]:
            raise ValueError(f"Spalte '{column}' fehlt in der Quelle.")
        return sum(
            (
                cls._decimal(row.get(column, ""), str(i))
                for i, row in enumerate(rows, 2)
            ),
            Decimal("0"),
        )

    @staticmethod
    def _normalise_decimal(raw: str) -> str:
        value = raw.strip()
        sign = value[:1] if value[:1] in {"+", "-"} else ""
        body = value[1:] if sign else value
        separators = {separator for separator in (".", ",") if separator in body}
        if not separators:
            return value

        if len(separators) == 2:
            decimal_separator = "," if body.rfind(",") > body.rfind(".") else "."
            grouping_separator = "." if decimal_separator == "," else ","
            if body.count(decimal_separator) != 1:
                raise ValueError("Mehrdeutiges Zahlenformat.")
            integer, fraction = body.rsplit(decimal_separator, 1)
            groups = integer.split(grouping_separator)
            if (
                not fraction.isdigit()
                or not groups
                or not groups[0].isdigit()
                or not 1 <= len(groups[0]) <= 3
                or any(not group.isdigit() or len(group) != 3 for group in groups[1:])
            ):
                raise ValueError("Mehrdeutiges Zahlenformat.")
            return f"{sign}{''.join(groups)}.{fraction}"

        separator = next(iter(separators))
        groups = body.split(separator)
        if not all(group.isdigit() for group in groups):
            raise ValueError("Mehrdeutiges Zahlenformat.")
        if len(groups) == 2:
            integer, fraction = groups
            # A single separator followed by exactly three digits can mean a
            # decimal or a thousands separator. Require an unambiguous source.
            if len(fraction) == 3 and integer != "0":
                raise ValueError("Mehrdeutiges Zahlenformat.")
            return f"{sign}{integer}.{fraction}"
        if (
            not groups
            or not 1 <= len(groups[0]) <= 3
            or any(len(group) != 3 for group in groups[1:])
        ):
            raise ValueError("Mehrdeutiges Zahlenformat.")
        return f"{sign}{''.join(groups)}"

    @staticmethod
    def _finite_decimal(raw: Any, label: str) -> Decimal:
        try:
            value = Decimal(str(raw))
        except (InvalidOperation, ValueError):
            raise ValueError(f"{label} muss eine endliche Zahl sein.") from None
        if not value.is_finite():
            raise ValueError(f"{label} muss eine endliche Zahl sein.")
        return value

    @classmethod
    def _decimal(cls, raw: str, account: str) -> Decimal:
        try:
            value = Decimal(cls._normalise_decimal(raw.strip()))
        except (InvalidOperation, ValueError):
            raise ValueError(f"Ungültige Zahl für Konto/Zeile '{account}'.") from None
        if not value.is_finite():
            raise ValueError(f"Ungültige Zahl für Konto/Zeile '{account}'.")
        return value

    @staticmethod
    def _format(value: Decimal) -> str:
        return format(value.normalize(), "f") if value else "0"

    @classmethod
    def _optional_decimal(cls, raw: str, account: str) -> Decimal | None:
        return None if not raw.strip() else cls._decimal(raw, account)

    @classmethod
    def _rows(
        cls, path: Path, required: list[str] | None = None
    ) -> list[dict[str, str]]:
        if path.suffix.lower() == ".csv":
            try:
                with path.open(encoding="utf-8-sig", newline="") as h:
                    sample = h.read(4096)
                    h.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                    except csv.Error:
                        dialect = csv.excel
                    reader = csv.DictReader(h, dialect=dialect)
                    rows = list(islice(reader, _MAX_SOURCE_ROWS + 1))
                    if reader.fieldnames and len(reader.fieldnames) > _MAX_SOURCE_COLUMNS:
                        raise ValueError("Eingabe enthält zu viele Spalten.")
            except (UnicodeDecodeError, csv.Error, OSError):
                raise ValueError("CSV muss UTF-8-kodiert sein.") from None
        else:
            cls._validate_xlsx_archive(path)
            try:
                from openpyxl import load_workbook
                from openpyxl.utils.exceptions import InvalidFileException
            except ImportError:
                raise ValueError("XLSX-Unterstützung ist nicht verfügbar.") from None
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
            except (
                BadZipFile,
                EOFError,
                InvalidFileException,
                KeyError,
                OSError,
                ParseError,
                ValueError,
            ):
                raise ValueError("XLSX konnte nicht gelesen werden.") from None
            try:
                sheet = wb.active
                allrows = list(
                    islice(sheet.iter_rows(values_only=True), _MAX_SOURCE_ROWS + 1)
                )
                if len(allrows) > _MAX_SOURCE_ROWS:
                    raise ValueError("Eingabe enthält zu viele Zeilen.")
                if any(len(row) > _MAX_SOURCE_COLUMNS for row in allrows):
                    raise ValueError("Eingabe enthält zu viele Spalten.")
                if allrows:
                    header = next(
                        (
                            i
                            for i, row in enumerate(allrows[:20])
                            if not required
                            or set(required)
                            <= set(str(x).strip() for x in row if x is not None)
                        ),
                        0,
                    )
                    names = [
                        str(x).strip() if x is not None else ""
                        for x in allrows[header]
                    ]
                    rows = [
                        {
                            names[i]: "" if x is None else str(x)
                            for i, x in enumerate(row[: len(names)])
                        }
                        for row in allrows[header + 1 :]
                        if any(x is not None for x in row)
                    ]
                else:
                    rows = []
            except (
                AttributeError,
                BadZipFile,
                EOFError,
                IndexError,
                KeyError,
                OSError,
                ParseError,
                TypeError,
                ValueError,
            ):
                raise ValueError("XLSX konnte nicht gelesen werden.") from None
            finally:
                wb.close()
        if len(rows) > _MAX_SOURCE_ROWS:
            raise ValueError("Eingabe enthält zu viele Zeilen.")
        if required and (not rows or not set(required) <= set(rows[0])):
            raise ValueError("Erforderliche Salden-Spalten fehlen.")
        return rows

    @staticmethod
    def _validate_xlsx_archive(path: Path) -> None:
        try:
            with ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > _MAX_XLSX_MEMBERS:
                    raise ValueError("XLSX enthält zu viele Archivdateien.")
                total_size = 0
                for member in members:
                    if member.file_size > _MAX_XLSX_MEMBER_BYTES:
                        raise ValueError("XLSX enthält eine zu große Archivdatei.")
                    total_size += member.file_size
                    if total_size > _MAX_XLSX_UNCOMPRESSED_BYTES:
                        raise ValueError("XLSX entpackt zu groß.")
        except (BadZipFile, LargeZipFile, OSError):
            raise ValueError("XLSX konnte nicht gelesen werden.") from None

    @classmethod
    def _metadata(cls, path: Path) -> tuple[int, list[str], list[str]]:
        rows = cls._rows(path)
        sheets = ["CSV"]
        if path.suffix.lower() == ".xlsx":
            cls._validate_xlsx_archive(path)
            try:
                from openpyxl import load_workbook
                from openpyxl.utils.exceptions import InvalidFileException
            except ImportError:
                raise ValueError("XLSX-Unterstützung ist nicht verfügbar.") from None
            try:
                workbook = load_workbook(path, read_only=True, data_only=True)
            except (
                BadZipFile,
                EOFError,
                InvalidFileException,
                KeyError,
                OSError,
                ParseError,
                ValueError,
            ):
                raise ValueError("XLSX konnte nicht gelesen werden.") from None
            try:
                sheets = list(workbook.sheetnames)
            except (
                AttributeError,
                BadZipFile,
                EOFError,
                IndexError,
                KeyError,
                OSError,
                ParseError,
                TypeError,
                ValueError,
            ):
                raise ValueError("XLSX konnte nicht gelesen werden.") from None
            finally:
                workbook.close()
        return (
            len(rows),
            list(rows[0]) if rows else [],
            sheets,
        )
